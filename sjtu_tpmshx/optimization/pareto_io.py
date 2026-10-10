"""Read legacy decision tables and export native optimization observations."""
import csv
import json
from pathlib import Path

import numpy as np

from sjtu_tpmshx.io.file_set import staged_files


def export_pareto_data(report, path):
    """Export the plotted native objectives, keeping archive IDs and run status."""
    path = Path(path)
    if path.suffix.lower() not in ('.csv', '.xlsx'):
        raise ValueError('Pareto data export requires .csv or .xlsx')
    history = report['history']
    front = set(report['pareto_indices'])
    rows = []
    for index, candidate in enumerate(history):
        if candidate['status'] != 'completed':
            continue
        objectives = candidate['objectives']
        pressure, heat = (float(objectives[key])
                          for key in ('pressure_ratio', 'heat_gain_percent'))
        if not np.isfinite([pressure, heat]).all():
            raise ValueError(f'Design {candidate["index"]} has nonfinite objectives')
        rows.append(dict(history_index=index, design_id=candidate['index'],
                         pressure_ratio=pressure, heat_gain_percent=heat,
                         is_pareto=index in front, directory=candidate['directory'],
                         run_status=report['status']))
    if not rows:
        raise ValueError('No completed, usable optimization observations to export')
    if not front <= {row['history_index'] for row in rows}:
        raise ValueError('Pareto indices must identify completed, usable observations')
    columns = list(rows[0])
    pareto_rows = sorted((row for row in rows if row['is_pareto']),
                         key=lambda row: row['pressure_ratio'])
    path.parent.mkdir(parents=True, exist_ok=True)
    with staged_files([path]) as stage:
        target = stage / path.name
        if path.suffix.lower() == '.csv':
            with target.open('w', newline='', encoding='utf-8-sig') as output:
                writer = csv.DictWriter(output, fieldnames=columns)
                writer.writeheader()
                writer.writerows(rows)
        else:
            from openpyxl import Workbook
            workbook = Workbook()
            try:
                sheet = workbook.active
                sheet.title = 'all_designs'
                front_sheet = workbook.create_sheet('pareto_front')
                for sheet, values in ((sheet, rows), (front_sheet, pareto_rows)):
                    sheet.append(columns)
                    for row in values:
                        sheet.append([row[key] for key in columns])
                metadata = workbook.create_sheet('metadata')
                metadata.append(['key', 'value'])
                for key in ('status', 'method', 'dimension', 'field_spec', 'objective_directions'):
                    value = report.get(key)
                    metadata.append([key, json.dumps(value, ensure_ascii=False)
                                     if isinstance(value, dict) else value])
                metadata.append(['x_axis', 'pressure_ratio [1]'])
                metadata.append(['y_axis', 'heat_gain_percent [%]'])
                workbook.save(target)
            finally:
                workbook.close()
    return dict(n_usable=len(rows), n_pareto=len(pareto_rows))


def read_pareto_csv(path, decision_dim_expected=None):
    with open(path, newline='', encoding='utf-8-sig') as source:
        reader = csv.DictReader(source)
        header = reader.fieldnames or []
        dimension = len(header) - 2 if decision_dim_expected is None else decision_dim_expected
        columns = [f'x{i}' for i in range(dimension)] + ['Q_W_per_m', 'dP_Pa']
        if dimension < 1 or len(header) != len(columns) or set(header) != set(columns):
            raise ValueError(f'Pareto CSV must contain x0..x{dimension - 1}, '
                             'Q_W_per_m and dP_Pa for the configured field layout')
        rows = []
        for index, row in enumerate(reader):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f'Pareto row {index} does not match the CSV columns')
            try:
                values = np.asarray([row[name] for name in columns], dtype=float)
            except ValueError as exc:
                raise ValueError(f'Pareto row {index} contains nonnumeric values') from exc
            if not np.all(np.isfinite(values)):
                raise ValueError(f'Pareto row {index} contains nonfinite values')
            rows.append(values)
    return np.asarray(rows, dtype=float).reshape(-1, len(columns))
