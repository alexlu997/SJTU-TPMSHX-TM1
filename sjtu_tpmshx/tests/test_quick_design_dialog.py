"""Construction tests for build_quick_design_dialog.

Verifies:
  - All contract attributes exist on the returned QDialog instance.
  - The run button calls run_quick_design.
  - Switching combo_qd_mode between "auto" and "fixed" toggles group visibility.

Run from repo root:
    python -m pytest sjtu_tpmshx/tests/test_quick_design_dialog.py -v
"""
from unittest.mock import patch
import pytest
from sjtu_tpmshx.ui.quick_design_panel import build_quick_design_dialog

CONTRACT = [
    "le_qd_file", "combo_qd_mode", "combo_qd_arr", "le_qd_rho",
    "le_qd_topo", "le_qd_l", "le_qd_t", "chk_qd_refine",
    "combo_qd_cell_topo", "le_qd_cell_l", "le_qd_cell_t",
    "_qd_table", "_qd_energy_table", "_qd_status",
]


def test_dialog_builds_with_contract_attrs():
    dlg = build_quick_design_dialog()
    for a in CONTRACT:
        assert hasattr(dlg, a), f"missing contract attr {a}"
    assert dlg.combo_qd_mode.currentData() in ("auto", "fixed")
    assert dlg.combo_qd_arr.currentData() in ("counter", "cross")
    from PySide6.QtWidgets import QTabWidget, QTableWidget
    tabs = dlg.findChild(QTabWidget)
    assert [tabs.tabText(i) for i in range(tabs.count())] == ['可行设计', '工况能量诊断']
    assert dlg._qd_table.columnCount() == 13
    assert dlg._qd_energy_table.editTriggers() == QTableWidget.EditTrigger.NoEditTriggers
    dlg.deleteLater()


def test_run_button_invokes_run_quick_design():
    dlg = build_quick_design_dialog()
    with patch("sjtu_tpmshx.ui.quick_design_panel.run_quick_design") as m:
        dlg._qd_run_btn.click()   # expose the run button as dlg._qd_run_btn
        assert m.called
    dlg.deleteLater()


def test_mode_toggle_switches_groups():
    dlg = build_quick_design_dialog()
    dlg.combo_qd_mode.setCurrentIndex(dlg.combo_qd_mode.findData("fixed"))
    assert dlg._qd_fixed_group.isVisibleTo(dlg) or not dlg._qd_auto_group.isVisibleTo(dlg)
    dlg.combo_qd_mode.setCurrentIndex(dlg.combo_qd_mode.findData("auto"))
    assert dlg._qd_auto_group.isVisibleTo(dlg) or not dlg._qd_fixed_group.isVisibleTo(dlg)
    dlg.deleteLater()


@pytest.mark.parametrize('mode', ['auto', 'fixed'])
@pytest.mark.parametrize('arrangement', ['counter', 'cross'])
@pytest.mark.parametrize('properties', ['mean', 'const'])
def test_display_labels_do_not_change_backend_values(mode, arrangement, properties):
    from sjtu_tpmshx.ui.quick_design_panel import _gather_inputs

    dlg = build_quick_design_dialog()
    dlg.combo_qd_mode.setCurrentIndex(dlg.combo_qd_mode.findData(mode))
    dlg.combo_qd_arr.setCurrentIndex(dlg.combo_qd_arr.findData(arrangement))
    dlg.combo_qd_prop.setCurrentIndex(dlg.combo_qd_prop.findData(properties))
    for combo in (dlg.combo_qd_mode, dlg.combo_qd_arr, dlg.combo_qd_prop):
        assert combo.count() == 2
        combo.setItemText(combo.currentIndex(), '修改后的显示名称')
    params = _gather_inputs(dlg)
    assert (params['mode'], params['arrangement'], params['prop_model']) == (
        mode, arrangement, properties)
    assert dlg._qd_auto_group.isHidden() == (mode != 'auto')
    assert dlg._qd_fixed_group.isHidden() == (mode != 'fixed')
    dlg.deleteLater()


@pytest.mark.parametrize('alias', [False, True])
def test_export_preserves_accepted_input_after_file_draft_changes(tmp_path, monkeypatch, alias):
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QFileDialog, QPushButton
    from sjtu_tpmshx.design.cases import load_cases
    from sjtu_tpmshx.design.sizing import Design
    from sjtu_tpmshx.tests.design.test_cases import _make_xlsx
    from sjtu_tpmshx.tests.gui_worker_support import _wait_for

    source = tmp_path / 'input.xlsx'
    _make_xlsx(source)
    previous = source.read_bytes()
    target = source
    if alias:
        target = tmp_path / 'input-alias.xlsx'
        try:
            target.symlink_to(source)
        except OSError:
            pytest.skip('symlinks unavailable')
    calls = []

    def candidate(cases, *args, **kwargs):
        calls.append(len(cases))
        return Design(True, topo='Diamond', l=7., t=.5, V=.001)

    monkeypatch.setattr('sjtu_tpmshx.design.sizing.size_fixed_cell', candidate)
    # Actual buttons, QThread, input loader and XLSX writer; only numerical
    # sizing and the headless file chooser's selected path are substituted.
    dialog = build_quick_design_dialog()
    try:
        dialog.le_qd_file.setText(str(source))
        dialog.combo_qd_mode.setCurrentIndex(dialog.combo_qd_mode.findData('fixed'))
        dialog._qd_run_btn.click()
        _wait_for(lambda: dialog._qd_worker is None)
        assert calls == [2]
        assert dialog._qd_last['params']['file'] == str(source)
        dialog.le_qd_file.setText(str(tmp_path / 'next-input.xlsx'))
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: (str(target), ''))
        export = next(b for b in dialog.findChildren(QPushButton) if b.text() == '导出 xlsx')
        export.click()
        assert dialog._qd_status.text() == '导出失败: 导出文件不能覆盖输入工况文件'
        assert source.read_bytes() == previous
        assert len(load_cases(str(source))) == 2
        if alias:
            assert target.is_symlink()
        target = tmp_path / 'report.xlsx'
        export.click()
        assert '已导出' in dialog._qd_status.text()
        workbook = load_workbook(target, read_only=True)
        try:
            assert workbook.sheetnames == ['构型汇总', '工况明细']
            cells = list(workbook['构型汇总'].values)
            row = dict(zip(cells[0], cells[1]))
            assert row['拓扑'] == 'Diamond' and row['l_mm'] == 7.
        finally:
            workbook.close()
        assert source.read_bytes() == previous
    finally:
        if getattr(dialog, '_qd_worker', None) is not None:
            dialog._qd_worker.requestInterruption()
            _wait_for(lambda: dialog._qd_worker is None)
        dialog.close()
