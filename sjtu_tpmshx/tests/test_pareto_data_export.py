"""Native objective exports preserve plotted values and incomplete-run status."""
import csv
from copy import deepcopy

from openpyxl import load_workbook
import pytest

from sjtu_tpmshx.optimization.pareto_io import export_pareto_data


@pytest.fixture
def report():
    return dict(status='cancelled', method='sobol', dimension=3,
                field_spec={'n_ctrl_x': 3, 'n_ctrl_y': 2, 'n_ctrl_z': 1},
                history=[
                    dict(index=0, status='completed', directory='design_0000',
                         objectives=dict(pressure_ratio=.8, heat_gain_percent=-2.)),
                    dict(index=1, status='failed', directory='design_0001', objectives=None),
                    dict(index=2, status='completed', directory='design_0002',
                         objectives=dict(pressure_ratio=.7, heat_gain_percent=3.))],
                pareto_indices=[2])


def test_csv_and_workbook_reproduce_plot_and_keep_ids(report, tmp_path):
    before = deepcopy(report)
    csv_path = tmp_path / 'points.csv'
    assert export_pareto_data(report, csv_path) == dict(n_usable=2, n_pareto=1)
    with csv_path.open(encoding='utf-8-sig', newline='') as source:
        rows = list(csv.DictReader(source))
    assert [int(row['design_id']) for row in rows] == [0, 2]
    assert [float(row['heat_gain_percent']) for row in rows] == [-2., 3.]
    assert [float(row['pressure_ratio']) for row in rows] == [.8, .7]
    assert [row['is_pareto'] for row in rows] == ['False', 'True']
    assert all(row['run_status'] == 'cancelled' for row in rows)
    path = tmp_path / 'points.xlsx'
    export_pareto_data(report, path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ['all_designs', 'pareto_front', 'metadata']
        points = list(workbook['all_designs'].values)
        front = list(workbook['pareto_front'].values)
        assert points[1][:5] == (0, 0, .8, -2, False)
        assert points[2][:5] == (2, 2, .7, 3, True)
        assert front == [points[0], points[2]]
        assert dict(list(workbook['metadata'].values)[1:])['status'] == 'cancelled'
    finally:
        workbook.close()
    assert report == before


@pytest.mark.parametrize('damage', ['nonfinite', 'failed_front', 'empty'])
def test_invalid_observation_does_not_replace_an_export(report, tmp_path, damage):
    path = tmp_path / 'points.csv'
    path.write_text('previous export')
    if damage == 'nonfinite':
        report['history'][0]['objectives']['heat_gain_percent'] = float('nan')
    elif damage == 'failed_front':
        report['pareto_indices'] = [1]
    else:
        report['history'] = []
    with pytest.raises(ValueError):
        export_pareto_data(report, path)
    assert path.read_text() == 'previous export'


def test_failed_workbook_write_keeps_existing_file(report, tmp_path, monkeypatch):
    from openpyxl import Workbook
    path = tmp_path / 'points.xlsx'
    path.write_bytes(b'previous workbook')

    def fail(self, path):
        path.write_bytes(b'incomplete')
        raise OSError('disk full')

    monkeypatch.setattr(Workbook, 'save', fail)
    with pytest.raises(OSError, match='disk full'):
        export_pareto_data(report, path)
    assert path.read_bytes() == b'previous workbook'
