"""An incomplete design workbook must not replace the previous report."""
from openpyxl import Workbook, load_workbook
import pytest

from sjtu_tpmshx.design import report
from sjtu_tpmshx.design.sizing import Design


def _previous_report(path):
    workbook = Workbook()
    workbook.active.title = 'previous-report'
    workbook.active['A1'] = 'keep the previous report'
    workbook.save(path)
    workbook.close()
    return path.read_bytes()


@pytest.mark.parametrize('existing', [False, True])
def test_second_sheet_failure_preserves_target(tmp_path, monkeypatch, existing):
    path = tmp_path / 'report.xlsx'
    previous = _previous_report(path) if existing else None
    original = report.pd.DataFrame.to_excel
    failure = OSError('second sheet write failed')

    def fail_second(frame, writer, *args, **kwargs):
        if kwargs.get('sheet_name') == '工况明细':
            raise failure
        return original(frame, writer, *args, **kwargs)

    monkeypatch.setattr(report.pd.DataFrame, 'to_excel', fail_second)
    with pytest.raises(OSError) as error:
        report.write_xlsx(path, [Design(True, topo='Gyroid', l=6., t=.4)])
    assert error.value is failure
    if existing:
        assert path.read_bytes() == previous
    else:
        assert not path.exists()
    assert not list(tmp_path.glob('.tm1-publish-*'))


def test_success_replaces_previous_report_with_both_sheets(tmp_path):
    path = tmp_path / 'report.xlsx'
    _previous_report(path)
    design = Design(True, topo='Gyroid', l=6., t=.4, V=.01, percase=[dict(
        case=3, hot_fluid='air', cold_fluid='water', T_air_out=350., T_cold_out=310.,
        dP_hot_pa=100., dP_hot_frac=.001, dP_cold_pa=200., dP_cold_frac=.002,
        Q_W=1250., Re_hot=1200., Re_cold=1500., run_status={'converged': True})])
    assert report.write_xlsx(str(path), [design]) == (1, 1, 1)
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ['构型汇总', '工况明细']
        summary = list(workbook['构型汇总'].values)
        assert summary[1][summary[0].index('可行')] == '是'
        details = list(workbook['工况明细'].values)
        row = dict(zip(details[0], details[1]))
        assert row['工况'] == 3
        assert row['换热量_kW'] == 1.25
        assert row['冷侧绝对压损_Pa'] == 200.
    finally:
        workbook.close()
    assert not list(tmp_path.glob('.tm1-publish-*'))
