"""The native backend completes real batches and a bounded design search."""
from dataclasses import replace
import json

import numpy as np
import pytest

from sjtu_tpmshx.domain.compute_config import FluidConfig
from sjtu_tpmshx.io.metrics_io import load_metrics
from sjtu_tpmshx.io.result_io import load_result
from sjtu_tpmshx.optimization.multi_condition import evaluate_condition_batch
from sjtu_tpmshx.optimization.multi_condition_optimizer import run_multi_condition_optimization
from sjtu_tpmshx.postprocess.api import evaluate
from sjtu_tpmshx.preprocess.api import prepare_case
from sjtu_tpmshx.preprocess.inlet_flow import total_inlet_mass_capacity
from sjtu_tpmshx.tests.native import test_cpp_full_3d as native_3d
from sjtu_tpmshx.tests.native.test_native_execution import _config
from sjtu_tpmshx.tests.gui_io_support import win as win

native_path = native_3d.native_path


def test_gui_co2_mass_input_save_reopen_geometry_and_export(win, native_path, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from sjtu_tpmshx.preprocess.api import prepare_inlet_mass_capacities
    from sjtu_tpmshx.tests.test_co2_model import co2_config
    from sjtu_tpmshx.tests.integration_tm1.test_public_gui import apply_config
    from sjtu_tpmshx.tests.gui_worker_support import _wait_for
    from sjtu_tpmshx.ui.window_config import config_from_window

    errors = []
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[1:]))
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: errors.append(args[1:]))
    def accept_grid(dialog):
        assert 'Grid preflight' in dialog.text(), dialog.text()
        return QMessageBox.StandardButton.Yes
    monkeypatch.setattr(QMessageBox, 'exec', accept_grid)
    config = co2_config(3)
    target = prepare_inlet_mass_capacities(config)['A'] * config.fluid_A.u_mps
    config.fluid_A.mass_flow_kg_s = target
    apply_config(win, config)
    native = native_3d.control(native_path)
    win.run_control = replace(win.run_control, native_library=native.native_library,
                              native_table_directory=native.native_table_directory)
    assert win.run_control.backend == 'cpp'
    saved = tmp_path/'co2-config.json'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: (str(saved), 'JSON'))
    win.save_config()
    assert saved.is_file(), errors
    win.combo_fluidA.setCurrentIndex(2)
    win.combo_inlet_modeA.setCurrentIndex(0)
    assert win._load_config_path(saved), errors
    reopened = config_from_window(win, strict=True)
    assert reopened.fluid_A.type == 'co2' and reopened.fluid_A.mass_flow_kg_s == target
    finished = []
    win.compute.finished.connect(finished.append)
    try:
        for factor in (1., 2.):
            depth = config.geometry.Lz_m * factor
            win.le_Lz.setText(str(depth))
            for side in 'AB':
                for end in ('in', 'out'):
                    getattr(win, f'le_pipe{side}_{end}_z_ctr').setText(str(depth/2))
                    getattr(win, f'le_pipe{side}_{end}_z_w').setText(str(depth))
            win.run_calculation()
            _wait_for(win.compute.is_idle, timeout=180)
            assert len(finished) == int(factor), errors
            result = finished[-1]
            assert result.converged and result.metadata['backend_id'] == 'cpp'
            assert result.metadata['co2']['pressure_multiplier'] == 2.5
            assert result.metadata['co2']['nu_multiplier'] == 1.28
            inlet = result.metadata['inlet_inputs']['A']
            assert inlet['mass_flow_kg_s'] == target
            assert inlet['u_mps'] == pytest.approx(config.fluid_A.u_mps/factor, rel=1e-12)
            assert float(win.le_uA.text()) == pytest.approx(inlet['u_mps'], rel=1e-10)
            assert abs(result.residuals['mass_imbalance_rel_A']) <= 1e-6
            output = tmp_path/f'co2-{int(factor)}.csv'
            monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a: (str(output), 'CSV'))
            win._export_results()
            assert output.is_file() and not errors
    finally:
        win.compute.finished.disconnect(finished.append)
        if not win.compute.is_idle():
            win.compute.cancel()
            _wait_for(win.compute.is_idle, timeout=30)


def _conditions(dimension, fluid='air'):
    config = replace(_config(dimension, True),
        fluid_B=FluidConfig(type='water', u_mps=.3, T_in_K=300., P_in_Pa=2e6))
    config = replace(config, geometry=replace(config.geometry, Lz_m=.03))
    if dimension == 3:
        config = replace(config, **{'bc_' + side: replace(getattr(config, 'bc_' + side),
            in_w=.03, out_w=.03, in_z_w=.03, out_z_w=.03) for side in 'AB'})
    conditions = []
    for index in range(2):
        cfg = replace(config, fluid_A=replace(config.fluid_A, type=fluid,
            u_mps=1.5 if fluid == 'sco2' else 12., T_in_K=500.+10*index,
            P_in_Pa=12e6 if fluid == 'sco2' else config.fluid_A.P_in_Pa))
        case = prepare_case(cfg, case_id=f'input-{index}')
        flows = [total_inlet_mass_capacity(case.design_fields, case.parameters, case.grid, side)
                 * getattr(cfg, 'fluid_' + side).u_mps for side in 'AB']
        conditions.append((f'condition-{index}', cfg, *flows))

    return conditions


@pytest.mark.parametrize('dimension', [2, 3])
@pytest.mark.parametrize('fluid', ['air', 'sco2'])
def test_selected_backend_completes_batch_and_optimization(native_path, tmp_path, dimension, fluid):
    conditions = _conditions(dimension, fluid)
    control = replace(native_3d.control(native_path), cpu_threads=4)
    workers = 2
    directory = tmp_path / 'cpp'
    batch = evaluate_condition_batch(conditions, output_dir=directory / 'batch', control=control, workers=workers)
    assert batch['status'] == 'completed', batch['conditions']
    assert batch['backend'] == 'cpp' and len(batch['results']) == 2
    for row, (_, field, metrics) in zip(batch['conditions'], batch['results']):
        restored = load_result(directory / 'batch' / row['result_file'])
        assert restored.backend_id == field.backend_id == 'cpp'
        assert restored.run_status['converged']
        assert row['qualification']['design']['eligible'] is (None if fluid == 'sco2' else True)
        assert row['qualification']['optimization']['passed'] is True
        assert load_metrics(directory / 'batch' / row['metrics_file']) == metrics
        assert evaluate(restored).metrics == metrics.metrics
        for name in ('Ta', 'Tb', 'Ts'):
            np.testing.assert_array_equal(restored.fields[name], field.fields[name])

    updates = []
    search = run_multi_condition_optimization(conditions, output_dir=directory / 'search',
        method='sobol', n_init=1, n_iter=1, seed=7, control=control, workers=workers,
        field_spec={'L_bounds': [6.8, 7.2], 't_bounds': [.55, .6]}, on_update=updates.append)
    assert search['status'] == 'completed', search
    assert search['backend'] == 'cpp'
    assert search['n_evaluated'] == search['n_usable'] == 2
    assert search['pareto_indices']
    assert updates[-1]['report'] == search
    assert [row['condition_id'] for row in updates[-1]['conditions']] == [row[0] for row in conditions]
    assert all(row['status'] == 'completed' for row in updates[-1]['conditions'])
    assert {'prepare', 'save_case', 'qualification'} <= {
        row['stage'] for update in updates for row in update['conditions']}
    for entry in (search['baseline'], *search['history']):
        batch_dir = directory / 'search' / entry['directory']
        record = json.loads((batch_dir / 'batch.json').read_text())
        assert record['backend'] == 'cpp' and record['status'] == 'completed'
        for row in record['conditions']:
            assert row['qualification']['design']['eligible'] is (None if fluid == 'sco2' else True)
            assert row['qualification']['optimization']['passed'] is True
            field = load_result(batch_dir / row['result_file'])
            assert field.backend_id == 'cpp' and field.run_status['converged']


@pytest.mark.parametrize('dimension', [2, 3])
def test_native_mass_search_preserves_total_units_constraints_and_exports(native_path, tmp_path, dimension):
    import csv
    from openpyxl import load_workbook
    from sjtu_tpmshx.optimization.multi_condition_optimizer import resume_multi_condition_optimization
    from sjtu_tpmshx.optimization.export_ntop_csv import export_study_design
    from sjtu_tpmshx.optimization.pareto_io import export_pareto_data

    conditions = _conditions(dimension)
    requirements = dict(solid_density_kg_m3=8000., conditions={name: dict(
        Q_min_W=1., dP_A_max_Pa=1e8, dP_B_max_Pa=1e8) for name, *_ in conditions})
    control = replace(native_3d.control(native_path), cpu_threads=1)
    directory = tmp_path / 'search'
    report = run_multi_condition_optimization(conditions, output_dir=directory,
        method='sobol', n_init=1, n_iter=1, seed=7, control=control, workers=1,
        field_spec={'L_bounds': [7., 7.], 't_bounds': [.55, .6]}, mass_requirements=requirements)
    assert report['status'] == 'completed', report
    assert report['n_feasible'] == report['n_usable'] == 2
    for row in report['history']:
        assert row['x_decision'][:len(row['x_decision'])//2] == [7.]*(len(row['x_decision'])//2)
        batch_dir = directory / row['directory']
        batch = json.loads((batch_dir / 'batch.json').read_text())
        field = load_result(batch_dir / batch['conditions'][0]['result_file'])
        depth = .03 if dimension == 2 else 1.
        volume = field.grid['dx'][:, None] * field.grid['dy'][None, :]
        if dimension == 3:
            volume = volume[:, :, None] * field.grid['dz'][None, None, :]
        expected = np.sum((1-field.metadata['design_fields']['eps_arr'])*volume)*8000.*depth
        assert row['objectives']['core_mass_kg'] == pytest.approx(expected)
        for case in batch['conditions']:
            metrics = load_metrics(batch_dir / case['metrics_file']).metrics
            check = row['constraint_checks'][case['condition_id']]
            assert check['Q_W'] == -metrics['Q_B'].value*depth
            assert check['dP_A_Pa'] == metrics['dP_A'].value
            assert check['dP_B_Pa'] == metrics['dP_B'].value
            assert check['feasible'] and max(check['residuals']) <= 0
    moved = directory.rename(tmp_path / 'moved')
    resumed = resume_multi_condition_optimization(moved / 'optimization.json',
        output_dir=tmp_path / 'resumed', control=control, workers=1)
    assert resumed['history'] == report['history']
    assert resumed['mass_requirements'] == requirements
    for suffix in ('.csv', '.xlsx'):
        path = tmp_path / ('results'+suffix)
        assert export_pareto_data(resumed, path) == dict(n_usable=2, n_pareto=1)
        if suffix == '.csv':
            with path.open(encoding='utf-8-sig') as source:
                rows = list(csv.DictReader(source))
            assert json.loads(rows[0]['mass_requirements']) == requirements
            assert float(rows[0]['core_mass_kg']) == report['history'][0]['objectives']['core_mass_kg']
        else:
            book = load_workbook(path, read_only=True)
            try:
                assert book['all_designs'].max_row == 3 and book['pareto_front'].max_row == 2
                assert json.loads(dict(book['metadata'].values)['mass_requirements']) == requirements
            finally:
                book.close()
    selected, = resumed['pareto_indices']
    export_study_design(resumed, selected, tmp_path / 'geometry', (3,)*dimension)
    lengths = np.genfromtxt(tmp_path / 'geometry/Lfield.csv', delimiter=',', skip_header=1)[:, -1]
    np.testing.assert_array_equal(lengths, 7.)
    metadata = json.loads((tmp_path / 'geometry/provenance.json').read_text())
    if dimension == 2:
        assert metadata['source']['physical_depth_m'] == .03
    assert metadata['source']['mass_requirements'] == requirements
    assert metadata['source']['constraint_checks'] == resumed['history'][selected]['constraint_checks']


@pytest.mark.parametrize('dimension', [2, 3])
def test_converged_outside_range_batch_retains_diagnostics(native_path, tmp_path, dimension):
    config = replace(_config(dimension, True),
        fluid_B=FluidConfig(type='water', u_mps=.2, T_in_K=300., P_in_Pa=2e6))
    config = replace(config, geometry=replace(config.geometry, Lz_m=.03))
    case = prepare_case(config, case_id='outside-range')
    flows = [total_inlet_mass_capacity(case.design_fields, case.parameters, case.grid, side)
             * getattr(config, 'fluid_' + side).u_mps for side in 'AB']
    control = native_3d.control(native_path)
    batch = evaluate_condition_batch([('outside', config, *flows)],
        output_dir=tmp_path / 'batch', control=control)
    row, = batch['conditions']
    assert batch['status'] == row['status'] == 'failed' and batch['objectives'] is None
    assert row['stage'] == 'qualification'
    field = load_result(tmp_path / 'batch' / row['result_file'])
    metrics = load_metrics(tmp_path / 'batch' / row['metrics_file'])
    assert field.run_status['converged']
    assert metrics.qualification['physical']['passed'] is True
    assert metrics.qualification['metrics']['passed'] is True
    assert metrics.qualification['applicability']['passed'] is False
    assert metrics.qualification['design']['eligible'] is False
    assert evaluate(field).qualification == metrics.qualification



def _wait_at_native_checkpoint(index, case_file, directory, library, threads, checkpoint_file=None):
    from time import monotonic, sleep
    from sjtu_tpmshx.optimization import multi_condition as batch
    from sjtu_tpmshx.workflows import recovery

    solve, save = recovery.run_case, recovery.save_solver_checkpoint
    def solve_with_inner_checkpoints(*args, **kwargs):
        kwargs['checkpoint_due'] = lambda: True
        return solve(*args, **kwargs)
    def save_and_wait(*args):
        save(*args)
        saved = args[0]
        state = saved.native_state
        if int(state['state.phase'][0]) != 6:
            return
        key = ('state.flow.value.0.iteration.value.result.iterations'
               if saved.state_format == 'full_2d_continuation_v1'
               else 'state.flow_iteration.value.iterations')
        if int(state[key][0]) < 2:
            return
        batch._batch_events.put((index, 'iteration', ('checkpoint-ready',)))
        deadline = monotonic()+30.
        while not batch._batch_cancel.is_set():
            if monotonic() > deadline:
                raise TimeoutError('The parent did not cancel the native test worker.')
            sleep(.01)
    recovery.run_case = solve_with_inner_checkpoints
    recovery.save_solver_checkpoint = save_and_wait
    return batch._run_native_condition(index, case_file, directory, library, threads, checkpoint_file)


@pytest.mark.parametrize('dimension', [2, 3])
def test_native_condition_workers_resume_actual_states_after_relocation(
        native_path, tmp_path, monkeypatch, dimension):
    from sjtu_tpmshx.domain.cancellation import CancelledError
    from sjtu_tpmshx.optimization import multi_condition as batch
    from sjtu_tpmshx.tests.native.test_run_recovery import same_result
    from sjtu_tpmshx.io.solver_checkpoint_io import load_solver_checkpoint

    conditions = _conditions(dimension)
    parallel = replace(native_3d.control(native_path), cpu_threads=4)
    if parallel.native_cpu_threads() < 2:
        pytest.skip('Two available CPUs are required for the real worker test.')
    threads = parallel.native_cpu_threads() // 2
    control = replace(parallel, cpu_threads=threads)
    expected = batch.evaluate_condition_batch(conditions, output_dir=tmp_path / 'expected', control=control)
    assert expected['status'] == 'completed', expected['conditions']
    native_worker = batch._run_native_condition
    monkeypatch.setattr(batch, '_run_native_condition', _wait_at_native_checkpoint)
    cancelled = []
    def iteration(message):
        if message == 'checkpoint-ready':
            cancelled.append(True)
    source = tmp_path / 'interrupted'
    with pytest.raises(CancelledError):
        batch.evaluate_condition_batch(conditions, output_dir=source, workers=2,
            control=replace(parallel, iteration=iteration, cancel_check=lambda: bool(cancelled)))
    record = json.loads((source / 'batch.json').read_text())
    assert record['status'] == 'cancelled' and record['objectives'] is None
    checkpoints = [load_solver_checkpoint(source / row['checkpoint_file'])
                   for row in record['conditions'] if row.get('checkpoint_file')]
    key = ('state.flow.value.0.iteration.value.result.iterations' if dimension == 2
           else 'state.flow_iteration.value.iterations')
    assert any(int(saved.native_state['state.phase'][0]) == 6
               and int(saved.native_state[key][0]) >= 2 for saved in checkpoints)
    moved = source.rename(tmp_path / 'moved')
    before = {path.relative_to(moved): path.read_bytes() for path in moved.rglob('*')
              if path.is_file() and path.suffix in ('.json', '.h5', '.yaml')}
    monkeypatch.setattr(batch, '_run_native_condition', native_worker)
    target = tmp_path / 'continued'
    actual = batch.resume_condition_batch(moved / 'batch.json', output_dir=target,
                                          workers=2, control=parallel)
    assert actual['status'] == 'completed', actual['conditions']
    assert any(row['recovery_mode'] == 'checkpoint' for row in actual['conditions'])
    for row, previous, returned, reference in zip(actual['conditions'], record['conditions'],
                                                  actual['results'], expected['results']):
        assert row['case_id'] == previous['case_id'] == returned[1].case_id
        assert row['cpu_threads'] == threads
        assert row['condition_id'] == returned[0] == reference[0]
        same_result(returned[1], reference[1], returned[2], reference[2])
        same_result(load_result(target / row['result_file']), reference[1],
                    load_metrics(target / row['metrics_file']), reference[2])
    assert {path.relative_to(moved): path.read_bytes() for path in moved.rglob('*')
            if path.is_file() and path.suffix in ('.json', '.h5', '.yaml')} == before


@pytest.mark.parametrize('dimension', [2, 3])
def test_gui_resumes_sou_and_exports_the_saved_result(win, native_path, tmp_path, monkeypatch, dimension):
    import csv
    import threading
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from sjtu_tpmshx.domain.paused_run import PausedRun
    from sjtu_tpmshx.tests.gui_worker_support import _wait_for
    from sjtu_tpmshx.tests.native.test_run_recovery import same_result
    from sjtu_tpmshx.workflows.compute import compute
    from sjtu_tpmshx.workflows.recovery import compute_run

    monkeypatch.setenv('TPMSHX_DISABLE_3D_PANEL', '1')
    errors = []
    for name in ('warning', 'critical'):
        monkeypatch.setattr(QMessageBox, name, lambda *args: errors.append(args[1:]))
    config = _config(dimension, True)
    config = replace(config, solver=replace(config.solver, enthalpy_algorithm='temperature_sou'))
    control = replace(native_3d.control(native_path), cpu_threads=1)
    case, expected, expected_metrics = compute(config, case_id='gui-sou', control=control)
    assert expected.run_status['converged']
    pause = threading.Event()
    source = tmp_path / 'paused'
    with pytest.raises(PausedRun):
        compute_run(config, source, case_id=case.case_id, checkpoint_interval_s=3600.,
            control=replace(control, iteration=lambda _: pause.set()), pause_check=pause.is_set)
    original = {name: (source / name).read_bytes() for name in ('case.h5', 'checkpoint.h5', 'run.json')}
    win.cache.clear()
    win.run_control = control
    win.le_Nx.setText('99')
    finished = []
    win.compute.finished.connect(finished.append)
    try:
        win._resume_saved_run(source / 'checkpoint.h5')
        _wait_for(win.compute.is_idle, timeout=180)
    finally:
        win.compute.finished.disconnect(finished.append)
        if not win.compute.is_idle():
            win.compute.cancel()
            _wait_for(win.compute.is_idle, timeout=30)
    assert len(finished) == 1 and not errors
    directory, = tmp_path.glob('run-*')
    actual = load_result(directory / 'results.h5')
    metrics = load_metrics(directory / 'metrics.json')
    same_result(actual, expected, metrics, expected_metrics)
    assert actual.boundary_fluxes['true_h']['energy_algorithm'] == 'temperature_sou'
    display = win.cache.get_result(f'{dimension}d')
    assert display.metadata['source_result_id'] == actual.result_id
    assert display.Q_W == metrics.metrics['Q'].value
    assert display.fields['Ta'].shape == expected.fields['Ta'].shape
    assert win.le_Nx.text() == '99'
    assert {name: (source / name).read_bytes() for name in original} == original
    output = tmp_path / 'sou.csv'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(output), 'CSV'))
    win._export_results()
    with output.open(encoding='utf-8-sig') as stream:
        rows = list(csv.reader(stream))
    assert [f'Q [{"W" if dimension == 3 else "W/m"}]', f'{display.Q_W:.4f}'] in rows
    assert not errors


@pytest.mark.parametrize('dimension', [2, 3])
def test_sou_batch_and_search_preserve_original_range_rejection(native_path, tmp_path, dimension):
    config = _config(dimension, True)
    config = replace(config, geometry=replace(config.geometry, Lz_m=.03),
                     solver=replace(config.solver, enthalpy_algorithm='temperature_sou'))
    case = prepare_case(config, case_id='sou-outside-range')
    flows = [total_inlet_mass_capacity(case.design_fields, case.parameters, case.grid, side)
             * getattr(config, 'fluid_' + side).u_mps for side in 'AB']
    conditions = [('outside', config, *flows)]
    control = replace(native_3d.control(native_path), cpu_threads=1)
    directory = tmp_path / 'batch'
    batch = evaluate_condition_batch(conditions, output_dir=directory, control=control, workers=1)
    row, = batch['conditions']
    assert batch['status'] == row['status'] == 'failed' and batch['objectives'] is None
    assert row['stage'] == 'qualification'
    field = load_result(directory / row['result_file'])
    metrics = load_metrics(directory / row['metrics_file'])
    assert field.run_status['converged'] and metrics.qualification['physical']['passed']
    assert field.metadata['thermal_mode'] == 'conservative_energy'
    assert field.boundary_fluxes['true_h']['energy_algorithm'] == 'temperature_sou'
    assert metrics.qualification['applicability']['passed'] is False
    assert metrics.qualification['design']['eligible'] is False
    assert evaluate(field).qualification == metrics.qualification
    search = run_multi_condition_optimization(conditions, output_dir=tmp_path / 'search',
        method='sobol', n_init=1, n_iter=0, seed=7, control=control, workers=1,
        field_spec={'L_bounds': [6.8, 7.2], 't_bounds': [.55, .6]})
    assert search['status'] == search['baseline']['status'] == 'failed'
    assert search['n_evaluated'] == search['n_usable'] == 0
    assert search['history'] == search['pareto_indices'] == []
    baseline_dir = tmp_path / 'search' / search['baseline']['directory']
    record = json.loads((baseline_dir / 'batch.json').read_text())
    assert record['objectives'] is None
    row, = record['conditions']
    assert row['stage'] == 'qualification'
    baseline = load_result(baseline_dir / row['result_file'])
    assert baseline.boundary_fluxes['true_h']['energy_algorithm'] == 'temperature_sou'
    assert row['qualification']['design']['eligible'] is False
